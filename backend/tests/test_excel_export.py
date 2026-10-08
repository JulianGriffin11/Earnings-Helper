"""Excel export workbook shape, without a database."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from app.services.excel_export import (
    DEBRIEF_HEADERS,
    DebriefExportRow,
    YoyExportRow,
    YOY_HEADERS,
    build_workbook,
    load_ticker_export,
    workbook_filename,
)


def test_view_sql_flattens_json() -> None:
    import importlib.util
    from pathlib import Path

    path = (
        Path(__file__).resolve().parents[1]
        / "alembic"
        / "versions"
        / "002_analyst_views.py"
    )
    spec = importlib.util.spec_from_file_location("analyst_views_migration", path)
    assert spec is not None and spec.loader is not None
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)

    yoy_sql = migration.ANALYST_YOY_SQL
    debrief_sql = migration.ANALYST_DEBRIEF_SQL

    assert "CREATE VIEW analyst_yoy" in yoy_sql
    assert "jsonb_array_elements(r.yoy_data->'metrics')" in yoy_sql
    assert "current_amount" in yoy_sql
    assert "pct_change" in yoy_sql
    assert "CREATE VIEW analyst_debrief" in debrief_sql
    assert "key_takeaways" in debrief_sql
    assert "items_to_watch" in debrief_sql
    assert "period_type = 'quarterly'" in debrief_sql
    assert "chr(10)" in debrief_sql


def test_workbook_filename_strips_unsafe_characters() -> None:
    assert workbook_filename("aapl") == "AAPL-earnings.xlsx"
    assert workbook_filename("brk.b") == "BRK.B-earnings.xlsx"
    assert workbook_filename("../aapl") == "AAPL-earnings.xlsx"


def test_workbook_sheets_match_the_analyst_tables() -> None:
    content = build_workbook(
        [
            YoyExportRow(
                ticker="AAPL",
                company="Apple Inc.",
                period_type="quarterly",
                period_end=date(2025, 6, 28),
                prior_period_end=date(2024, 6, 29),
                filing_date="2025-08-01",
                metric="Revenue",
                current_amount=Decimal("94036000000"),
                prior_amount=Decimal("85777000000"),
                dollar_change=Decimal("8259000000"),
                pct_change=Decimal("9.63"),
            ),
            YoyExportRow(
                ticker="AAPL",
                company="Apple Inc.",
                period_type="quarterly",
                period_end=date(2025, 6, 28),
                prior_period_end=date(2024, 6, 29),
                filing_date="2025-08-01",
                metric="Net Income",
                current_amount=None,
                prior_amount=None,
                dollar_change=None,
                pct_change=None,
            ),
        ],
        [
            DebriefExportRow(
                ticker="AAPL",
                filing_date="2025-08-01",
                headline="Revenue rose.",
                overall_assessment="mixed",
                margin_analysis="Gross profit grew faster than revenue.",
                takeaways="Quarterly revenue rose.\nFull-year net income fell.",
                items_to_watch="Operating expenses grew faster than annual revenue.",
            )
        ],
    )

    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["YoY", "Debrief"]

    yoy = workbook["YoY"]
    assert [cell.value for cell in yoy[1]] == YOY_HEADERS
    assert yoy["A2"].value == "AAPL"
    assert yoy["C2"].value == "quarterly"
    assert yoy["D2"].value.date() == date(2025, 6, 28)
    assert yoy["D2"].number_format == "YYYY-MM-DD"
    assert yoy["G2"].value == "Revenue"
    assert yoy["H2"].value == 94036000000
    assert yoy["H2"].number_format == "#,##0.00"
    assert yoy["K2"].value == pytest.approx(9.63)
    assert yoy["K2"].number_format == '0.00"%"'
    assert yoy["H3"].value is None
    assert "YoY" in yoy.tables
    assert yoy.tables["YoY"].ref == "A1:K3"

    debrief = workbook["Debrief"]
    assert [cell.value for cell in debrief[1]] == DEBRIEF_HEADERS
    assert debrief["C2"].value == "Revenue rose."
    assert debrief["D2"].value == "mixed"
    assert debrief["F2"].value == "Quarterly revenue rose.\nFull-year net income fell."
    assert debrief["F2"].alignment.wrap_text is True
    assert "Debrief" in debrief.tables
    assert debrief.tables["Debrief"].ref == "A1:G2"


def test_debrief_sheet_without_rows_has_headers_only() -> None:
    content = build_workbook(
        [
            YoyExportRow(
                ticker="AAPL",
                company="Apple Inc.",
                period_type="annual",
                period_end=None,
                prior_period_end=None,
                filing_date="unknown",
                metric="Revenue",
                current_amount=1,
                prior_amount=1,
                dollar_change=0,
                pct_change=0,
            )
        ],
        [],
    )
    debrief = load_workbook(BytesIO(content))["Debrief"]
    assert [cell.value for cell in debrief[1]] == DEBRIEF_HEADERS
    assert debrief["A2"].value is None
    assert "Debrief" not in debrief.tables


def test_missing_ticker_raises() -> None:
    db = Mock()
    empty = Mock()
    empty.mappings.return_value.all.return_value = []
    db.execute.return_value = empty

    with pytest.raises(LookupError, match="No cached report"):
        load_ticker_export(db, "ZZZZ")
    db.execute.assert_called_once()
