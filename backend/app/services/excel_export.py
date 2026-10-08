"""Build an Excel workbook from the analyst views for one ticker."""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import text
from sqlalchemy.orm import Session

YOY_HEADERS = [
    "Ticker",
    "Company",
    "Period Type",
    "Period End",
    "Prior Period End",
    "Filing Date",
    "Metric",
    "Current",
    "Prior",
    "Dollar Change",
    "Percent Change",
]

DEBRIEF_HEADERS = [
    "Ticker",
    "Filing Date",
    "Headline",
    "Overall Assessment",
    "Margin Note",
    "Takeaways",
    "Items to Watch",
]

MONEY_FORMAT = "#,##0.00"
# pct_change is stored in percentage points (9.63 means 9.63%).
PERCENT_FORMAT = '0.00"%"'
DATE_FORMAT = "YYYY-MM-DD"

YOY_MONEY_COLUMNS = {8, 9, 10}
YOY_PERCENT_COLUMN = 11
YOY_DATE_COLUMNS = {4, 5}
DEBRIEF_WRAP_COLUMNS = {5, 6, 7}

YOY_WIDTHS = [12, 28, 14, 14, 18, 14, 22, 18, 18, 16, 16]
DEBRIEF_WIDTHS = [12, 14, 48, 20, 42, 48, 42]

YOY_EXPORT_SQL = text(
    """
    SELECT
        ticker,
        company,
        period_type,
        period_end,
        prior_period_end,
        filing_date,
        metric,
        current_amount,
        prior_amount,
        dollar_change,
        pct_change
    FROM analyst_yoy
    WHERE upper(ticker) = upper(:ticker)
    ORDER BY
        filing_date DESC,
        CASE period_type
            WHEN 'quarterly' THEN 0
            WHEN 'annual' THEN 1
            ELSE 2
        END,
        metric_order
    """
)

DEBRIEF_EXPORT_SQL = text(
    """
    SELECT
        ticker,
        filing_date,
        headline,
        overall_assessment,
        margin_analysis,
        takeaways,
        items_to_watch
    FROM analyst_debrief
    WHERE upper(ticker) = upper(:ticker)
    ORDER BY filing_date DESC
    """
)

Number = int | float | Decimal | None


class YoyExportRow:
    def __init__(
        self,
        *,
        ticker: str,
        company: str,
        period_type: str,
        period_end: date | None,
        prior_period_end: date | None,
        filing_date: str,
        metric: str,
        current_amount: Number,
        prior_amount: Number,
        dollar_change: Number,
        pct_change: Number,
    ) -> None:
        self.ticker = ticker
        self.company = company
        self.period_type = period_type
        self.period_end = period_end
        self.prior_period_end = prior_period_end
        self.filing_date = filing_date
        self.metric = metric
        self.current_amount = current_amount
        self.prior_amount = prior_amount
        self.dollar_change = dollar_change
        self.pct_change = pct_change


class DebriefExportRow:
    def __init__(
        self,
        *,
        ticker: str,
        filing_date: str,
        headline: str | None,
        overall_assessment: str | None,
        margin_analysis: str | None,
        takeaways: str | None,
        items_to_watch: str | None,
    ) -> None:
        self.ticker = ticker
        self.filing_date = filing_date
        self.headline = headline
        self.overall_assessment = overall_assessment
        self.margin_analysis = margin_analysis
        self.takeaways = takeaways
        self.items_to_watch = items_to_watch


def workbook_filename(ticker: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]", "", ticker).strip(".-").upper()
    return f"{safe or 'REPORT'}-earnings.xlsx"


def load_ticker_export(
    db: Session,
    ticker: str,
) -> tuple[list[YoyExportRow], list[DebriefExportRow]]:
    yoy_rows = [
        _yoy_row(dict(row))
        for row in db.execute(YOY_EXPORT_SQL, {"ticker": ticker}).mappings().all()
    ]
    if not yoy_rows:
        raise LookupError("No cached report for this ticker")

    debrief_rows = [
        _debrief_row(dict(row))
        for row in db.execute(DEBRIEF_EXPORT_SQL, {"ticker": ticker}).mappings().all()
    ]
    return yoy_rows, debrief_rows


def build_ticker_workbook(db: Session, ticker: str) -> tuple[bytes, str]:
    yoy_rows, debrief_rows = load_ticker_export(db, ticker)
    return build_workbook(yoy_rows, debrief_rows), workbook_filename(ticker)


def build_workbook(
    yoy_rows: list[YoyExportRow],
    debrief_rows: list[DebriefExportRow],
) -> bytes:
    workbook = Workbook()
    yoy_sheet = workbook.active
    yoy_sheet.title = "YoY"
    _write_yoy_sheet(yoy_sheet, yoy_rows)

    debrief_sheet = workbook.create_sheet("Debrief")
    _write_debrief_sheet(debrief_sheet, debrief_rows)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_yoy_sheet(sheet: Worksheet, rows: list[YoyExportRow]) -> None:
    sheet.append(YOY_HEADERS)
    for row in rows:
        sheet.append(
            [
                row.ticker,
                row.company,
                row.period_type,
                row.period_end,
                row.prior_period_end,
                row.filing_date,
                row.metric,
                _number(row.current_amount),
                _number(row.prior_amount),
                _number(row.dollar_change),
                _number(row.pct_change),
            ]
        )
        excel_row = sheet.max_row
        for column in YOY_MONEY_COLUMNS:
            sheet.cell(excel_row, column).number_format = MONEY_FORMAT
        sheet.cell(excel_row, YOY_PERCENT_COLUMN).number_format = PERCENT_FORMAT
        for column in YOY_DATE_COLUMNS:
            cell = sheet.cell(excel_row, column)
            if isinstance(cell.value, date):
                cell.number_format = DATE_FORMAT

    _add_table(sheet, "YoY", len(rows), len(YOY_HEADERS))
    _size_columns(sheet, YOY_WIDTHS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = None


def _write_debrief_sheet(sheet: Worksheet, rows: list[DebriefExportRow]) -> None:
    sheet.append(DEBRIEF_HEADERS)
    for row in rows:
        values = [
            row.ticker,
            row.filing_date,
            row.headline,
            row.overall_assessment,
            row.margin_analysis,
            row.takeaways,
            row.items_to_watch,
        ]
        sheet.append(values)
        excel_row = sheet.max_row
        line_count = 1
        for column in DEBRIEF_WRAP_COLUMNS:
            cell = sheet.cell(excel_row, column)
            if isinstance(cell.value, str) and "\n" in cell.value:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                line_count = max(line_count, cell.value.count("\n") + 1)
        if line_count > 1:
            sheet.row_dimensions[excel_row].height = 15 * line_count

    _add_table(sheet, "Debrief", len(rows), len(DEBRIEF_HEADERS))
    _size_columns(sheet, DEBRIEF_WIDTHS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = None


def _add_table(sheet: Worksheet, name: str, row_count: int, col_count: int) -> None:
    if row_count < 1:
        return
    ref = f"A1:{get_column_letter(col_count)}{row_count + 1}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showColumnStripes=False,
        showFirstColumn=False,
        showLastColumn=False,
    )
    sheet.add_table(table)


def _size_columns(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _number(value: Number) -> float | None:
    if value is None:
        return None
    return float(value)


def _as_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    return str(value)


def _yoy_row(row: dict[str, Any]) -> YoyExportRow:
    return YoyExportRow(
        ticker=str(row["ticker"]),
        company=str(row["company"]),
        period_type=str(row["period_type"]),
        period_end=_as_date(row["period_end"]),
        prior_period_end=_as_date(row["prior_period_end"]),
        filing_date=str(row["filing_date"]),
        metric=str(row["metric"]),
        current_amount=row["current_amount"],
        prior_amount=row["prior_amount"],
        dollar_change=row["dollar_change"],
        pct_change=row["pct_change"],
    )


def _debrief_row(row: dict[str, Any]) -> DebriefExportRow:
    return DebriefExportRow(
        ticker=str(row["ticker"]),
        filing_date=str(row["filing_date"]),
        headline=_text(row["headline"]),
        overall_assessment=_text(row["overall_assessment"]),
        margin_analysis=_text(row["margin_analysis"]),
        takeaways=_text(row["takeaways"]),
        items_to_watch=_text(row["items_to_watch"]),
    )
